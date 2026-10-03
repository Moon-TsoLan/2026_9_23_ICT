"""Native readers for formats the application machine can read without a model.

PDFs no longer live here: they are parsed page-run by the parsing service, so nothing in
this module can resurrect the old `attachments-md` tree or its page-range guessing.
"""

from __future__ import annotations

from pathlib import Path


def read_docx(path: Path) -> list[dict]:
    import docx

    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = []
    for index, table in enumerate(document.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        tables.append({"table_index": index, "headers": rows[0] if rows else [], "rows": rows[1:]})
        for row in rows:
            parts.append(" | ".join(row))
    text = "\n".join(parts)
    return [{"page_no": 1, "text": text, "tables": tables, "chars": len(text), "source": "docx"}]


def read_xlsx(path: Path) -> list[dict]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages = []
    for index, sheet in enumerate(book.worksheets, start=1):
        rows = [["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(row)]
        text = "\n".join("\t".join(row) for row in rows[:400])
        pages.append({"page_no": index, "text": text,
                      "tables": [{"table_index": 0, "headers": rows[0] if rows else [], "rows": rows[1:]}],
                      "chars": len(text), "source": "xlsx"})
    book.close()
    return pages


def read_native(path: Path) -> list[dict]:
    """docx and spreadsheets keep their own structure; no rasterising, no model."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise ValueError("unsupported native document " + suffix)
